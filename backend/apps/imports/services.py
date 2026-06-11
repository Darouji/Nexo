import csv
import io
from pathlib import Path

from django.contrib.contenttypes.models import ContentType
from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from apps.core.models import AuditOrigin
from apps.core.utils import normalize_text
from apps.inventory.models import Categoria, KardexMovementType, MovimientoKardex, Producto
from apps.inventory.services import register_kardex_movement

from .models import ImportBatch, ImportBatchStatus, ImportRow, ImportRowStatus, ImportTarget


PRODUCT_HEADER_ALIASES = {
    "categoria": "categoria",
    "category": "categoria",
    "categoría": "categoria",
    "modelo": "modelo",
    "producto": "modelo",
    "model": "modelo",
    "nombre": "modelo",
    "detalle": "detalle",
    "descripcion": "detalle",
    "descripción": "detalle",
    "codigo_sap": "codigo_sap",
    "código_sap": "codigo_sap",
    "codigo sap": "codigo_sap",
    "código sap": "codigo_sap",
    "sap": "codigo_sap",
    "stock_inicial": "stock_inicial",
    "stock inicial": "stock_inicial",
    "stock": "stock_inicial",
    "stock_critico": "stock_critico",
    "stock crítico": "stock_critico",
    "stock critico": "stock_critico",
    "en_camino": "en_camino",
    "en camino": "en_camino",
    "controla_stock": "controla_stock",
    "controla stock": "controla_stock",
}


def get_operator(user):
    if user is None:
        return None

    if not getattr(user, "is_authenticated", False):
        return None

    return user


def normalize_header(value):
    return normalize_text(str(value or "").replace("_", " "))


def canonical_header(value):
    normalized = normalize_header(value)
    return PRODUCT_HEADER_ALIASES.get(normalized, normalized.replace(" ", "_"))


def as_text(value):
    if value is None:
        return ""

    return str(value).strip()


def as_int(value, *, default=0):
    if value is None or value == "":
        return default

    if isinstance(value, bool):
        return int(value)

    try:
        return int(float(value))
    except (TypeError, ValueError):
        raise ValidationError(f"Valor numérico inválido: {value}")


def as_bool(value, *, default=True):
    if value is None or value == "":
        return default

    if isinstance(value, bool):
        return value

    normalized = normalize_text(value)

    if normalized in ["si", "sí", "s", "true", "1", "x", "yes", "y"]:
        return True

    if normalized in ["no", "n", "false", "0"]:
        return False

    raise ValidationError(f"Valor booleano inválido: {value}")


def read_import_rows(batch):
    extension = Path(batch.original_filename or batch.file.name).suffix.lower()

    if extension == ".xlsx":
        return read_xlsx_rows(batch)

    if extension == ".csv":
        return read_csv_rows(batch)

    raise ValidationError({"file": "Formato de archivo no soportado."})


def read_xlsx_rows(batch):
    batch.file.open("rb")
    workbook = None

    try:
        workbook = load_workbook(batch.file, read_only=True, data_only=True)
        worksheet = workbook.active

        row_iterator = worksheet.iter_rows(values_only=True)
        header_row = next(row_iterator, None)

        if header_row is None:
            return []

        headers = [canonical_header(value) for value in header_row]
        output = []

        for index, row in enumerate(row_iterator, start=2):
            if row_is_empty(row):
                continue

            raw = {}

            for column_index, header in enumerate(headers):
                if not header:
                    continue

                raw[header] = row[column_index] if column_index < len(row) else ""

            output.append(
                {
                    "row_number": index,
                    "raw_data": raw,
                }
            )

        return output

    finally:
        if workbook is not None:
            workbook.close()

        batch.file.close()


def read_csv_rows(batch):
    batch.file.open("rb")

    try:
        raw_bytes = batch.file.read()

        try:
            text = raw_bytes.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw_bytes.decode("latin-1")

        reader = csv.DictReader(io.StringIO(text))
        output = []

        for index, row in enumerate(reader, start=2):
            normalized_row = {
                canonical_header(key): value
                for key, value in row.items()
                if key
            }

            if not any(as_text(value) for value in normalized_row.values()):
                continue

            output.append(
                {
                    "row_number": index,
                    "raw_data": normalized_row,
                }
            )

        return output
    finally:
        batch.file.close()


def row_is_empty(row):
    return not any(value not in [None, ""] for value in row)


def normalize_product_row(raw_data):
    categoria = as_text(raw_data.get("categoria"))
    modelo = as_text(raw_data.get("modelo"))
    detalle = as_text(raw_data.get("detalle"))
    codigo_sap = as_text(raw_data.get("codigo_sap"))

    stock_inicial = as_int(raw_data.get("stock_inicial"), default=0)
    stock_critico = as_int(raw_data.get("stock_critico"), default=0)
    en_camino = as_int(raw_data.get("en_camino"), default=0)
    controla_stock = as_bool(raw_data.get("controla_stock"), default=True)

    return {
        "categoria": categoria,
        "categoria_normalized": normalize_text(categoria),
        "modelo": modelo,
        "modelo_normalized": normalize_text(modelo),
        "detalle": detalle,
        "codigo_sap": codigo_sap,
        "codigo_sap_normalized": normalize_text(codigo_sap),
        "stock_inicial": stock_inicial,
        "stock_critico": stock_critico,
        "en_camino": en_camino,
        "controla_stock": controla_stock,
    }


def validate_product_row(normalized_data, seen_model_keys, seen_sap_keys):
    messages = []
    status = ImportRowStatus.VALIDO

    if not normalized_data["categoria"]:
        messages.append("La categoría es obligatoria.")

    if not normalized_data["modelo"]:
        messages.append("El modelo es obligatorio.")

    if normalized_data["stock_inicial"] < 0:
        messages.append("El stock inicial no puede ser negativo.")

    if normalized_data["stock_critico"] < 0:
        messages.append("El stock crítico no puede ser negativo.")

    if normalized_data["en_camino"] < 0:
        messages.append("El campo en_camino no puede ser negativo.")

    model_key = (
        normalized_data["categoria_normalized"],
        normalized_data["modelo_normalized"],
    )
    sap_key = normalized_data["codigo_sap_normalized"]

    if model_key in seen_model_keys:
        messages.append("Producto duplicado dentro del mismo archivo por categoría + modelo.")
        status = ImportRowStatus.DUPLICADO

    if sap_key and sap_key in seen_sap_keys:
        messages.append("Producto duplicado dentro del mismo archivo por código SAP.")
        status = ImportRowStatus.DUPLICADO

    if messages and status != ImportRowStatus.DUPLICADO:
        status = ImportRowStatus.ERROR

    if status == ImportRowStatus.VALIDO:
        seen_model_keys.add(model_key)

        if sap_key:
            seen_sap_keys.add(sap_key)

        existing_by_model = Producto.objects.filter(
            categoria__name_normalized=normalized_data["categoria_normalized"],
            modelo_normalized=normalized_data["modelo_normalized"],
        ).exists()

        existing_by_sap = False

        if sap_key:
            existing_by_sap = Producto.objects.filter(
                codigo_sap_normalized=sap_key,
            ).exists()

        if existing_by_model or existing_by_sap:
            status = ImportRowStatus.REQUIERE_REVISION
            messages.append(
                "Ya existe un producto similar. No se importará automáticamente para evitar sobrescritura silenciosa."
            )

    return status, messages


def prevalidate_import_batch(*, batch, user=None):
    operator = get_operator(user)

    if batch.batch_status == ImportBatchStatus.PROCESADO:
        raise ValidationError(
            {"batch_status": "No se puede prevalidar un lote ya procesado."}
        )

    if batch.target != ImportTarget.PRODUCTOS:
        raise ValidationError({"target": "Destino de importación no soportado todavía."})

    parsed_rows = read_import_rows(batch)

    with transaction.atomic():
        batch.rows.all().delete()

        seen_model_keys = set()
        seen_sap_keys = set()
        import_rows = []

        for parsed_row in parsed_rows:
            raw_data = parsed_row["raw_data"]

            try:
                normalized_data = normalize_product_row(raw_data)
                row_status, messages = validate_product_row(
                    normalized_data,
                    seen_model_keys,
                    seen_sap_keys,
                )
            except ValidationError as error:
                normalized_data = {}
                row_status = ImportRowStatus.ERROR
                messages = [str(error)]

            import_rows.append(
                ImportRow(
                    batch=batch,
                    row_number=parsed_row["row_number"],
                    row_status=row_status,
                    raw_data=raw_data,
                    normalized_data=normalized_data,
                    messages=messages,
                    created_by=operator,
                    updated_by=operator,
                )
            )

        ImportRow.objects.bulk_create(import_rows)

        batch.batch_status = ImportBatchStatus.PREVALIDADO
        batch.updated_by = operator
        batch.save(update_fields=["batch_status", "updated_by", "updated_at"])
        batch.update_counters()

        return batch


def get_or_create_category(category_name, user=None):
    normalized_name = normalize_text(category_name)
    category = Categoria.objects.filter(name_normalized=normalized_name).first()

    if category:
        return category

    return Categoria.objects.create(
        name=category_name,
        created_by=get_operator(user),
        updated_by=get_operator(user),
    )


def process_import_batch(*, batch, user=None):
    operator = get_operator(user)

    if batch.batch_status == ImportBatchStatus.PROCESADO:
        raise ValidationError({"batch_status": "El lote ya fue procesado."})

    if batch.batch_status != ImportBatchStatus.PREVALIDADO:
        batch = prevalidate_import_batch(batch=batch, user=operator)

    valid_rows = batch.rows.filter(row_status=ImportRowStatus.VALIDO).order_by("row_number")

    product_content_type = ContentType.objects.get_for_model(Producto)

    for row in valid_rows:
        try:
            with transaction.atomic():
                data = row.normalized_data

                category = get_or_create_category(data["categoria"], user=operator)

                existing_by_model = Producto.objects.filter(
                    categoria=category,
                    modelo_normalized=data["modelo_normalized"],
                ).exists()

                existing_by_sap = False

                if data["codigo_sap_normalized"]:
                    existing_by_sap = Producto.objects.filter(
                        codigo_sap_normalized=data["codigo_sap_normalized"],
                    ).exists()

                if existing_by_model or existing_by_sap:
                    row.row_status = ImportRowStatus.REQUIERE_REVISION
                    row.messages = [
                        "El producto apareció existente al procesar. No se importó para evitar sobrescritura silenciosa."
                    ]
                    row.updated_by = operator
                    row.save(
                        update_fields=[
                            "row_status",
                            "messages",
                            "updated_by",
                            "updated_at",
                        ]
                    )
                    continue

                product = Producto.objects.create(
                    categoria=category,
                    modelo=data["modelo"],
                    detalle=data["detalle"],
                    codigo_sap=data["codigo_sap"],
                    stock_critico=data["stock_critico"],
                    en_camino=data["en_camino"],
                    controla_stock=data["controla_stock"],
                    created_by=operator,
                    updated_by=operator,
                )

                if product.controla_stock and data["stock_inicial"] > 0:
                    register_kardex_movement(
                        producto=product,
                        movement_type=KardexMovementType.ENTRADA_IMPORTACION,
                        cantidad=data["stock_inicial"],
                        motivo=f"Entrada por importación lote #{batch.pk}, fila {row.row_number}",
                        user=operator,
                        origin=AuditOrigin.IMPORTACION_EXCEL,
                        reference=batch,
                    )

                row.row_status = ImportRowStatus.IMPORTADO
                row.imported_content_type = product_content_type
                row.imported_object_id = str(product.pk)
                row.imported_repr = str(product)[:255]
                row.updated_by = operator
                row.save(
                    update_fields=[
                        "row_status",
                        "imported_content_type",
                        "imported_object_id",
                        "imported_repr",
                        "updated_by",
                        "updated_at",
                    ]
                )

        except Exception as error:
            row.row_status = ImportRowStatus.ERROR
            row.messages = [f"Error al procesar fila: {error}"]
            row.updated_by = operator
            row.save(
                update_fields=[
                    "row_status",
                    "messages",
                    "updated_by",
                    "updated_at",
                ]
            )

    batch.batch_status = ImportBatchStatus.PROCESADO
    batch.processed_by = operator
    batch.processed_at = timezone.now()
    batch.updated_by = operator
    batch.save(
        update_fields=[
            "batch_status",
            "processed_by",
            "processed_at",
            "updated_by",
            "updated_at",
        ]
    )
    batch.update_counters()

    return batch