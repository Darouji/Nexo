# Construcción de exportaciones de Inventario desde definiciones de grilla
# autorizadas. La vista puede respetar columnas visibles, mientras que el modo
# oficial mantiene un contrato estable y auditable definido por el backend.
import csv
from io import BytesIO, StringIO

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook

from apps.core.grid_services import get_effective_column_labels
from apps.core.models import AuditAction, AuditLog, AuditOrigin

from .grid_definitions import (
    KARDEX_GRID_DEFINITION,
    PENDING_RECEIPT_GRID_DEFINITION,
    PRODUCT_GRID_DEFINITION,
)


def neutralize_spreadsheet_formula(value):
    # Evita que texto controlado por usuarios sea interpretado como fórmula al
    # abrir CSV o XLSX en aplicaciones de hojas de cálculo.
    if value is None:
        return ""

    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"

    return value


def product_export_value(product, field_key):
    # Las claves técnicas de filtros pueden apuntar a relaciones o propiedades;
    # la exportación devuelve sus valores visibles y no las formas normalizadas.
    if field_key == "categoria__name":
        return product.categoria.name
    if field_key == "por_recibir_total":
        return int(getattr(product, "por_recibir_total", 0) or 0)
    if field_key == "origen":
        return product.get_origen_display()
    if field_key == "stock_status":
        return product.stock_status
    if field_key == "estado_producto_display":
        return product.estado_producto_display

    return getattr(product, field_key, "")


def resolve_export_columns(
    *,
    mode,
    requested_columns,
    definition=PRODUCT_GRID_DEFINITION,
):
    column_map = definition.column_map

    # El modo oficial ignora preferencias personales para producir siempre el
    # mismo conjunto de datos; el modo de vista respeta columnas solicitadas.
    if mode == "official":
        field_keys = list(definition.official_export_columns)
    else:
        field_keys = requested_columns or [
            column.key for column in definition.columns if column.exportable
        ]

    # El backend valida cada clave para impedir exportar campos no declarados o
    # columnas operativas como acciones de interfaz.
    invalid = [
        field_key
        for field_key in field_keys
        if field_key not in column_map or not column_map[field_key].exportable
    ]
    if invalid:
        from rest_framework.exceptions import ValidationError

        raise ValidationError(
            {"columns": f"Columnas no exportables: {', '.join(invalid)}."}
        )

    return field_keys


def build_product_export_response(*, queryset, mode, export_format, columns):
    definition = PRODUCT_GRID_DEFINITION
    # Sólo la exportación de vista adopta etiquetas personalizadas; la oficial
    # conserva las etiquetas de sistema del contrato de la grilla.
    labels = get_effective_column_labels(definition) if mode == "view" else {}

    headers = [
        neutralize_spreadsheet_formula(
            labels.get(field_key, definition.column_map[field_key].system_label)
        )
        for field_key in columns
    ]
    rows = [
        [
            neutralize_spreadsheet_formula(product_export_value(product, field_key))
            for field_key in columns
        ]
        for product in queryset
    ]

    filename = f"serkan_catalogo_productos_{mode}.{export_format}"

    if export_format == "csv":
        output = StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerows(rows)
        response = HttpResponse(
            # La marca UTF-8 favorece la apertura correcta de tildes y caracteres
            # chilenos en aplicaciones que no detectan automáticamente el encoding.
            "\ufeff" + output.getvalue(),
            content_type="text/csv; charset=utf-8",
        )
    else:
        # write_only limita el consumo de memoria al generar planillas con muchas filas.
        workbook = Workbook(write_only=True)
        worksheet = workbook.create_sheet("Catálogo")
        worksheet.append(headers)
        for row in rows:
            worksheet.append(row)

        output = BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
        )

    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def pending_receipt_export_value(pending_receipt, field_key):
    # Los nombres normalizados sirven para consultar y ordenar; el archivo debe
    # conservar los valores legibles registrados para producto, categoría y ubicación.
    if field_key == "producto__modelo_normalized":
        return pending_receipt.producto.modelo
    if field_key == "producto__categoria__name_normalized":
        return pending_receipt.producto.categoria.name
    if field_key == "producto__codigo_sap_normalized":
        return pending_receipt.producto.codigo_sap
    if field_key == "destination_location__name_normalized":
        return pending_receipt.destination_location.name
    if field_key == "cantidad_pendiente_orden":
        return pending_receipt.cantidad_pendiente
    if field_key == "estado":
        return pending_receipt.get_estado_display()
    if field_key == "created_by__username":
        if not pending_receipt.created_by:
            return ""
        return (
            pending_receipt.created_by.get_full_name().strip()
            or pending_receipt.created_by.username
        )
    if field_key == "created_at":
        return timezone.localtime(pending_receipt.created_at).strftime(
            "%d-%m-%Y %H:%M:%S"
        )

    return getattr(pending_receipt, field_key, "")


def build_pending_receipt_export_response(
    *,
    queryset,
    mode,
    export_format,
    columns,
):
    definition = PENDING_RECEIPT_GRID_DEFINITION
    labels = get_effective_column_labels(definition) if mode == "view" else {}
    headers = [
        neutralize_spreadsheet_formula(
            labels.get(field_key, definition.column_map[field_key].system_label)
        )
        for field_key in columns
    ]
    rows = [
        [
            neutralize_spreadsheet_formula(
                pending_receipt_export_value(pending_receipt, field_key)
            )
            for field_key in columns
        ]
        for pending_receipt in queryset
    ]
    filename = f"serkan_recepciones_fisicas_{mode}.{export_format}"

    if export_format == "csv":
        output = StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerows(rows)
        response = HttpResponse(
            "\ufeff" + output.getvalue(),
            content_type="text/csv; charset=utf-8",
        )
    else:
        workbook = Workbook(write_only=True)
        worksheet = workbook.create_sheet("Recepciones físicas")
        worksheet.append(headers)
        for row in rows:
            worksheet.append(row)

        output = BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
        )

    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def record_grid_export(
    *,
    request,
    definition,
    mode,
    export_format,
    columns,
    row_count,
    model_name,
    object_repr,
):
    # Exportar información constituye una acción auditable. Se registra el
    # contrato utilizado y la cantidad de filas, no el contenido completo del archivo.
    AuditLog.objects.create(
        actor=request.user,
        action=AuditAction.EXPORTAR,
        origin=AuditOrigin.API,
        module="inventory",
        model_name=model_name,
        object_repr=object_repr,
        ip_address=request.META.get("REMOTE_ADDR"),
        user_agent=request.META.get("HTTP_USER_AGENT", "")[:500],
        metadata={
            "request_id": getattr(request, "request_id", ""),
            "grid_key": definition.key,
            "mode": mode,
            "format": export_format,
            "columns": columns,
            "row_count": row_count,
        },
    )


def record_product_export(*, request, mode, export_format, columns, row_count):
    record_grid_export(
        request=request,
        definition=PRODUCT_GRID_DEFINITION,
        mode=mode,
        export_format=export_format,
        columns=columns,
        row_count=row_count,
        model_name="producto",
        object_repr="Catálogo de productos",
    )


def kardex_export_value(movement, field_key):
    if field_key == "created_at":
        return timezone.localtime(movement.created_at).strftime("%d-%m-%Y %H:%M:%S")
    # Los snapshots conservan la información histórica aunque el producto,
    # categoría o ubicación cambien de nombre posteriormente.
    if field_key == "producto_modelo_orden":
        return movement.producto_modelo_snapshot or movement.producto.modelo
    if field_key == "producto_codigo_sap_orden":
        return movement.producto_codigo_sap_snapshot or movement.producto.codigo_sap
    if field_key == "producto_categoria_orden":
        return movement.producto_categoria_snapshot or movement.producto.categoria.name
    if field_key == "movement_type":
        return movement.get_movement_type_display()
    if field_key == "origin":
        return movement.get_origin_display()
    if field_key == "ubicacion_nombre":
        if movement.location_name_snapshot:
            return movement.location_name_snapshot
        return movement.location.name if movement.location else "No informada históricamente"
    if field_key == "operador_nombre":
        if not movement.operador:
            return ""
        return movement.operador.get_full_name().strip() or movement.operador.username
    return getattr(movement, field_key, "")


def build_kardex_export_response(*, queryset, mode, export_format, columns):
    definition = KARDEX_GRID_DEFINITION
    labels = get_effective_column_labels(definition) if mode == "view" else {}
    headers = [
        neutralize_spreadsheet_formula(
            labels.get(field_key, definition.column_map[field_key].system_label)
        )
        for field_key in columns
    ]
    rows = [
        [
            neutralize_spreadsheet_formula(kardex_export_value(movement, field_key))
            for field_key in columns
        ]
        for movement in queryset
    ]
    filename = f"serkan_kardex_global_{mode}.{export_format}"

    if export_format == "csv":
        output = StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerows(rows)
        response = HttpResponse(
            "\ufeff" + output.getvalue(), content_type="text/csv; charset=utf-8"
        )
    else:
        workbook = Workbook(write_only=True)
        worksheet = workbook.create_sheet("Kardex global")
        worksheet.append(headers)
        for row in rows:
            worksheet.append(row)
        output = BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
