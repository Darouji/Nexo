import json
from io import BytesIO

from django.test import TestCase
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APIClient

from apps.accounts.models import User, UserRole
from apps.core.models import AuditAction, AuditLog

from .models import Categoria, InventoryLocation, Producto
from .services import register_manual_entry


class KardexGridApiTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            username="kardex-grid-admin",
            email="kardex-grid-admin@serkan.cl",
            password="test-pass-123",
            role=UserRole.ADMINISTRADOR,
        )
        self.reader = User.objects.create_user(
            username="kardex-grid-reader",
            email="kardex-grid-reader@serkan.cl",
            password="test-pass-123",
            role=UserRole.CONSULTA,
        )
        category = Categoria.objects.create(name="Categoría Kardex")
        self.product = Producto.objects.create(
            categoria=category,
            modelo="Alternador =24V",
            codigo_sap="KDX-001",
            controla_stock=True,
        )
        self.location = InventoryLocation.objects.create(
            name="Bodega Kardex",
            location_type="BODEGA",
        )
        self.movement = register_manual_entry(
            producto=self.product,
            cantidad=4,
            motivo="=Ingreso inicial",
            location=self.location,
            user=self.admin,
        )

    def test_filters_location_and_remote_grid_before_pagination(self):
        client = APIClient()
        client.force_authenticate(user=self.admin)
        response = client.get(
            "/api/inventory/kardex/",
            {
                "ubicacion": self.location.pk,
                "grid_filters": json.dumps(
                    [{"field": "cantidad", "operator": "greaterThan", "value": 3}]
                ),
            },
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["count"], 1)
        self.assertEqual(response.data["results"][0]["id"], self.movement.pk)

    def test_view_export_is_safe_and_audited(self):
        client = APIClient()
        client.force_authenticate(user=self.admin)
        response = client.get(
            "/api/inventory/kardex/export/",
            {
                "mode": "view",
                "file_format": "xlsx",
                "columns": "producto_modelo_orden,motivo,cantidad",
            },
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        rows = list(load_workbook(BytesIO(response.content), read_only=True).active.iter_rows(values_only=True))
        self.assertEqual(rows[1][0], "Alternador =24V")
        self.assertEqual(rows[1][1], "'=Ingreso inicial")
        self.assertEqual(rows[1][2], 4)
        audit = AuditLog.objects.get(action=AuditAction.EXPORTAR, model_name="movimientokardex")
        self.assertEqual(audit.metadata["row_count"], 1)

    def test_export_rejects_empty_filtered_results(self):
        client = APIClient()
        client.force_authenticate(user=self.admin)
        response = client.get(
            "/api/inventory/kardex/export/",
            {"mode": "view", "file_format": "xlsx", "search": "SIN-RESULTADOS"},
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("No hay resultados para exportar", str(response.data))
        self.assertFalse(
            AuditLog.objects.filter(
                action=AuditAction.EXPORTAR,
                model_name="movimientokardex",
            ).exists()
        )

    def test_reader_cannot_request_official_export(self):
        client = APIClient()
        client.force_authenticate(user=self.reader)
        response = client.get(
            "/api/inventory/kardex/export/",
            {"mode": "official", "file_format": "xlsx"},
        )
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

