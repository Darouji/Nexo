import json
from io import BytesIO

from django.test import TestCase
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APIClient

from apps.accounts.models import User, UserRole
from apps.core.models import AuditAction, AuditLog

from .models import Categoria, InventoryLocation, Producto
from .services import create_pending_receipt, receive_pending_receipt


class PendingReceiptGridApiTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            username="pending-grid-admin",
            email="pending-grid-admin@serkan.cl",
            password="test-pass-123",
            role=UserRole.ADMINISTRADOR,
        )
        self.reader = User.objects.create_user(
            username="pending-grid-reader",
            email="pending-grid-reader@serkan.cl",
            password="test-pass-123",
            role=UserRole.CONSULTA,
        )
        category = Categoria.objects.create(name="Categoría pendientes")
        self.product = Producto.objects.create(
            categoria=category,
            modelo="Sensor pendiente",
            codigo_sap="PEND-001",
            controla_stock=True,
        )
        self.location = InventoryLocation.objects.create(
            name="Bodega pendientes",
            location_type="BODEGA",
        )
        self.pending_receipt = create_pending_receipt(
            producto=self.product,
            cantidad=8,
            referencia="=OC-PELIGROSA",
            motivo="Ingreso esperado.",
            destination_location=self.location,
            user=self.admin,
        )

    def test_remote_grid_filter_uses_pending_quantity_annotation(self):
        client = APIClient()
        client.force_authenticate(user=self.admin)

        response = client.get(
            "/api/inventory/pendientes-por-recibir/",
            {
                "grid_filters": json.dumps(
                    [
                        {
                            "field": "cantidad_pendiente_orden",
                            "operator": "greaterThan",
                            "value": 7,
                        }
                    ]
                )
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["count"], 1)
        self.assertEqual(
            response.data["results"][0]["id"],
            self.pending_receipt.pk,
        )

    def test_view_export_neutralizes_formulas_and_records_audit(self):
        client = APIClient()
        client.force_authenticate(user=self.admin)

        response = client.get(
            "/api/inventory/pendientes-por-recibir/export/",
            {
                "mode": "view",
                "file_format": "xlsx",
                "columns": (
                    "producto__modelo_normalized,referencia,"
                    "cantidad_pendiente_orden"
                ),
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        self.assertEqual(rows[1][0], "Sensor pendiente")
        self.assertEqual(rows[1][1], "'=OC-PELIGROSA")
        self.assertEqual(rows[1][2], 8)

        audit = AuditLog.objects.get(
            action=AuditAction.EXPORTAR,
            model_name="pendienteporrecibir",
        )
        self.assertEqual(
            audit.metadata["grid_key"],
            "inventory.pending-receipts",
        )
        self.assertEqual(audit.metadata["row_count"], 1)

    def test_export_rejects_empty_filtered_results(self):
        client = APIClient()
        client.force_authenticate(user=self.admin)
        response = client.get(
            "/api/inventory/pendientes-por-recibir/export/",
            {"mode": "view", "file_format": "xlsx", "search": "SIN-RESULTADOS"},
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("No hay resultados para exportar", str(response.data))
        self.assertFalse(
            AuditLog.objects.filter(
                action=AuditAction.EXPORTAR,
                model_name="pendienteporrecibir",
            ).exists()
        )

    def test_official_export_requires_official_export_capability(self):
        reader_client = APIClient()
        reader_client.force_authenticate(user=self.reader)
        denied_response = reader_client.get(
            "/api/inventory/pendientes-por-recibir/export/",
            {"mode": "official", "file_format": "csv"},
        )

        admin_client = APIClient()
        admin_client.force_authenticate(user=self.admin)
        allowed_response = admin_client.get(
            "/api/inventory/pendientes-por-recibir/export/",
            {"mode": "official", "file_format": "csv"},
        )

        self.assertEqual(denied_response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(allowed_response.status_code, status.HTTP_200_OK)

    def test_export_rejects_unknown_columns(self):
        client = APIClient()
        client.force_authenticate(user=self.admin)

        response = client.get(
            "/api/inventory/pendientes-por-recibir/export/",
            {
                "mode": "view",
                "file_format": "xlsx",
                "columns": (
                    "producto__modelo_normalized,campo_inexistente"
                ),
            },
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_history_preserves_readable_product_snapshots(self):
        receive_pending_receipt(
            pending_receipt=self.pending_receipt,
            cantidad=3,
            motivo="Recepción física validada.",
            idempotency_key="pending-history-receive-001",
            user=self.admin,
        )
        self.product.refresh_from_db()
        self.product.modelo = "Sensor renombrado"
        self.product.save(update_fields=["modelo"])

        client = APIClient()
        client.force_authenticate(user=self.admin)
        response = client.get(
            (
                "/api/inventory/pendientes-por-recibir/"
                f"{self.pending_receipt.pk}/historial/"
            )
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data), 1)
        self.assertEqual(
            response.data[0]["product_model_snapshot"],
            "Sensor pendiente",
        )
        self.assertEqual(
            response.data[0]["destination_location_snapshot"],
            "Bodega pendientes",
        )
        self.assertEqual(response.data[0]["pending_before"], 8)
        self.assertEqual(response.data[0]["pending_after"], 5)
